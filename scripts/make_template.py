"""Genera il modello Excel di Cartina di Classe senza dipendenze runtime."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "assets" / "Modello_Planimetria_Classi.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)

BLACK = "FF000000"
NAVY = "FF003F8F"
LIGHT_BLUE = "FFE3EDFF"
WHITE = "FFFFFFFF"
GREY = "FFF3F6FA"
DARK_GREY = "FF46556B"
WARN = "FFFFF7DF"

FONT = "Verdana"
thin = Side(style="thin", color=BLACK)
medium = Side(style="medium", color=BLACK)
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---------------------------------------------------------------------------
# Istruzioni
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "Istruzioni"
ws.sheet_view.showGridLines = False
ws.merge_cells("A1:H2")
c = ws["A1"]
c.value = "Cartina di Classe — modello Excel"
c.font = Font(name=FONT, size=24, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=NAVY)
c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True, indent=1)
ws.row_dimensions[1].height = 34
ws.row_dimensions[2].height = 34

ws.merge_cells("A4:H4")
c = ws["A4"]
c.value = "Come si usa"
c.font = Font(name=FONT, size=17, bold=True, color=BLACK)
c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
c.alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[4].height = 28

steps = [
    (6, "1.", "Apri il foglio «Layout» e scrivi quanti banchi ci sono in ogni cella."),
    (7, "2.", "Apri il foglio «Alunni» e incolla i nomi, uno per riga."),
    (8, "3.", "Salva questo file, poi importalo nell'app con il pulsante «Importa Excel»."),
]
for row, num, text in steps:
    ws.cell(row=row, column=1, value=num).font = Font(name=FONT, size=15, bold=True, color=NAVY)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    c = ws.cell(row=row, column=2, value=text)
    c.font = Font(name=FONT, size=14, color=BLACK)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28

ws.merge_cells("A10:H10")
c = ws["A10"]
c.value = "La disposizione dei banchi"
c.font = Font(name=FONT, size=17, bold=True, color=BLACK)
c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
c.alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[10].height = 28

paragraphs = [
    (12, "La prima riga del foglio Layout è l'ultima fila, quella in fondo all'aula. L'ultima riga è la prima fila, vicino alla lavagna."),
    (13, "Ogni numero indica quanti banchi ci sono in quel gruppo. Scrivi 0 per lasciare uno spazio vuoto."),
    (14, "Puoi aggiungere altre righe e colonne nel foglio Layout, fino a 20 file e 12 colonne."),
]
for row, text in paragraphs:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=13, color=BLACK)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[row].height = 34

# Esempio visivo
ws.merge_cells("A16:H16")
c = ws["A16"]
c.value = "Esempio"
c.font = Font(name=FONT, size=15, bold=True, color=BLACK)
ws.row_dimensions[16].height = 24

example = [
    ["", "Sinistra", "Centro", "Destra"],
    ["Ultima fila (fondo)", 2, 3, 2],
    ["Fila 2", 2, 2, 2],
    ["Prima fila (lavagna)", 2, 2, 1],
]
for ri, row_vals in enumerate(example, start=17):
    for ci, val in enumerate(row_vals, start=1):
        cell = ws.cell(row=ri, column=ci, value=val)
        cell.font = Font(name=FONT, size=14, bold=(ri == 17), color=BLACK)
        cell.border = border_all
        cell.alignment = Alignment(horizontal="center" if ci > 1 else "left", vertical="center")
        if ri == 17:
            cell.fill = PatternFill("solid", fgColor=GREY)
    ws.row_dimensions[ri].height = 24

ws.merge_cells("A22:H22")
c = ws["A22"]
c.value = "Nell'esempio: 3 banchi al centro dell'ultima fila; 1 banco a destra della prima fila."
c.font = Font(name=FONT, size=12, italic=True, color=DARK_GREY)
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[22].height = 24

ws.merge_cells("A24:H24")
c = ws["A24"]
c.value = "I nomi degli alunni"
c.font = Font(name=FONT, size=17, bold=True, color=BLACK)
c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
c.alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[24].height = 28

name_paragraphs = [
    (26, "Nel foglio «Alunni» scrivi un nome per riga nella colonna Alunno. Puoi scrivere nome e cognome nella stessa cella."),
    (27, "Se vuoi, compila anche Fila, Colonna e Banco per assegnare un posto preciso. Se lasci vuoto, l'app dispone l'alunno automaticamente."),
    (28, "Fila 1 = prima fila, vicino alla lavagna. Colonna 1 = sinistra, vista dal docente. Banco 1 = primo banco del gruppo, da sinistra."),
    (29, "Se un alunno non ha un posto indicato, l'app lo mette nel primo banco libero."),
]
for row, text in name_paragraphs:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=13, color=BLACK)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[row].height = 34

ws.merge_cells("A31:H31")
c = ws["A31"]
c.value = "Se Excel non viene importato"
c.font = Font(name=FONT, size=17, bold=True, color=BLACK)
c.fill = PatternFill("solid", fgColor=WARN)
c.alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[31].height = 28

ws.merge_cells("A33:H33")
c = ws["A33"]
c.value = "Apri il file in Excel e salvalo come «CSV UTF-8», poi importa il CSV. In alternativa, incolla i nomi direttamente nella scheda Alunni dell'app."
c.font = Font(name=FONT, size=13, color=BLACK)
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[33].height = 34

ws.merge_cells("A35:H35")
c = ws["A35"]
c.value = "L'app non invia dati in rete e funziona offline."
c.font = Font(name=FONT, size=12, italic=True, color=DARK_GREY)
c.alignment = Alignment(vertical="center")

for col in "ABCDEFGH":
    ws.column_dimensions[col].width = 16 if col != "A" else 20
ws.column_dimensions["A"].width = 26
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)

# ---------------------------------------------------------------------------
# Layout
# ---------------------------------------------------------------------------
wl = wb.create_sheet("Layout")
wl.sheet_view.showGridLines = False
wl.merge_cells("A1:F1")
c = wl["A1"]
c.value = "Disposizione dei banchi"
c.font = Font(name=FONT, size=22, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=NAVY)
c.alignment = Alignment(vertical="center", indent=1)
wl.row_dimensions[1].height = 34

wl.merge_cells("A2:F2")
c = wl["A2"]
c.value = "La prima riga è l'ULTIMA fila (fondo aula). L'ultima riga è la PRIMA fila (vicino alla lavagna)."
c.font = Font(name=FONT, size=14, bold=True, color=BLACK)
c.alignment = Alignment(wrap_text=True, vertical="center")
wl.row_dimensions[2].height = 40

wl.merge_cells("A3:F3")
c = wl["A3"]
c.value = "Scrivi quanti banchi ci sono in ogni cella. Usa 0 per uno spazio vuoto."
c.font = Font(name=FONT, size=13, color=BLACK)
c.alignment = Alignment(wrap_text=True, vertical="center")
wl.row_dimensions[3].height = 28

headers = ["", "Colonna 1 (sinistra)", "Colonna 2 (centro)", "Colonna 3 (destra)"]
for ci, h in enumerate(headers, start=1):
    cell = wl.cell(row=4, column=ci, value=h)
    cell.font = Font(name=FONT, size=12, bold=True, color=WHITE if ci > 1 else BLACK)
    cell.fill = PatternFill("solid", fgColor=NAVY if ci > 1 else WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border_all
wl.row_dimensions[4].height = 34

row_labels = ["Ultima fila (fondo)", "Fila 2", "Prima fila (lavagna)"]
for ri, label in enumerate(row_labels, start=5):
    c = wl.cell(row=ri, column=1, value=label)
    c.font = Font(name=FONT, size=12, bold=True, color=BLACK)
    c.fill = PatternFill("solid", fgColor=GREY)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    c.border = border_all
    for ci in range(2, 5):
        cell = wl.cell(row=ri, column=ci, value=2)
        cell.font = Font(name=FONT, size=18, bold=True, color=BLACK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_all
        cell.fill = PatternFill("solid", fgColor=WHITE)
    wl.row_dimensions[ri].height = 38

wl.merge_cells("A10:F10")
c = wl["A10"]
c.value = "Per aggiungere file o colonne, continua sotto o a destra. Massimo 20 file e 12 colonne. I numeri devono essere interi da 0 a 12."
c.font = Font(name=FONT, size=12, color=DARK_GREY)
c.alignment = Alignment(wrap_text=True, vertical="top")
wl.row_dimensions[10].height = 30

for ci in range(1, 7):
    wl.column_dimensions[get_column_letter(ci)].width = 22
wl.column_dimensions["A"].width = 30
wl.freeze_panes = "B5"
wl.page_setup.orientation = "landscape"
wl.page_setup.fitToWidth = 1
wl.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------------------
# Alunni
# ---------------------------------------------------------------------------
wa = wb.create_sheet("Alunni")
wa.sheet_view.showGridLines = False
wa.merge_cells("A1:D1")
c = wa["A1"]
c.value = "Elenco degli alunni"
c.font = Font(name=FONT, size=22, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=NAVY)
c.alignment = Alignment(vertical="center", indent=1)
wa.row_dimensions[1].height = 34

headers = [
    ("Alunno", "Scrivi nome e cognome. È obbligatorio."),
    ("Fila (facoltativo)", "1 = prima fila, vicino alla lavagna. Lascia vuoto per la disposizione automatica."),
    ("Colonna (facoltativo)", "1 = sinistra, vista dal docente. Lascia vuoto per la disposizione automatica."),
    ("Banco (facoltativo)", "1 = primo banco del gruppo, da sinistra. Lascia vuoto per la disposizione automatica."),
]
for ci, (h, comment) in enumerate(headers, start=1):
    cell = wa.cell(row=3, column=ci, value=h)
    cell.font = Font(name=FONT, size=13, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border_all
    cell.comment = Comment(comment, "Cartina di Classe")
    wa.column_dimensions[get_column_letter(ci)].width = 26 if ci == 1 else 15
wa.row_dimensions[3].height = 36

dv_row = DataValidation(type="whole", operator="between", formula1="1", formula2="20",
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Fila non valida", error="Inserisci un numero da 1 a 20.",
                        promptTitle="Fila", prompt="1 = prima fila, vicino alla lavagna.")
dv_col = DataValidation(type="whole", operator="between", formula1="1", formula2="12",
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Colonna non valida", error="Inserisci un numero da 1 a 12.",
                        promptTitle="Colonna", prompt="1 = sinistra, vista dal docente.")
dv_desk = DataValidation(type="whole", operator="between", formula1="1", formula2="12",
                         allow_blank=True, showErrorMessage=True,
                         errorTitle="Banco non valido", error="Inserisci un numero da 1 a 12.",
                         promptTitle="Banco", prompt="1 = primo banco del gruppo, da sinistra.")
wa.add_data_validation(dv_row)
wa.add_data_validation(dv_col)
wa.add_data_validation(dv_desk)
dv_row.add("B4:B60")
dv_col.add("C4:C60")
dv_desk.add("D4:D60")

for r in range(4, 61):
    for ci in range(1, 5):
        cell = wa.cell(row=r, column=ci)
        cell.border = border_all
        cell.font = Font(name=FONT, size=14, color=BLACK)
        cell.alignment = Alignment(vertical="center", horizontal="center" if ci > 1 else "left")
    wa.row_dimensions[r].height = 26

wa.freeze_panes = "A4"
wa.auto_filter.ref = "A3:D60"
wa.page_setup.orientation = "landscape"
wa.page_setup.fitToWidth = 1
wa.sheet_properties.pageSetUpPr.fitToPage = True

# Ordine dei fogli
wb._sheets = [wb["Istruzioni"], wb["Layout"], wb["Alunni"]]

wb.properties.title = "Modello Planimetria Classi"
wb.properties.subject = "Cartina di Classe"
wb.properties.creator = "Cartina di Classe"
wb.properties.description = "Modello per importare disposizione e alunni nell'app Cartina di Classe."

wb.save(OUT)
print("Creato", OUT, OUT.stat().st_size, "byte")
